"""Generate the TatRaksha pitch deck (12 slides, 16:9).

    uv run --with python-pptx python docs/pitch/build_deck.py

Writes docs/pitch/TatRaksha-pitch.pptx. Screenshots come from docs/pitch/img/ (captured from the running app).
Structured around the evaluation criteria: AI/technical 25 %, problem-solution fit 20 %, depth & reach across
India 20 %, deployability & scalability 20 %, impact 15 %.
"""
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

HERE = Path(__file__).resolve().parent
IMG = HERE / "img"
OUT = HERE / "TatRaksha-pitch.pptx"

NAVY = RGBColor(0x0B, 0x25, 0x45)
TEAL = RGBColor(0x0E, 0x7C, 0x86)
ORANGE = RGBColor(0xE0, 0x6C, 0x1F)
INK = RGBColor(0x1F, 0x29, 0x37)
MUTED = RGBColor(0x5B, 0x67, 0x78)
LIGHT = RGBColor(0xF1, 0xF5, 0xF9)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "Calibri"

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
W, H = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]


def box(slide, x, y, w, h, fill=None, line=None, shape=MSO_SHAPE.RECTANGLE):
    s = slide.shapes.add_shape(shape, x, y, w, h)
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(1)
    s.shadow.inherit = False
    return s


def text(slide, x, y, w, h, runs, size=18, color=INK, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
         spacing=1.1):
    """runs: str, or list of paragraphs; each paragraph a str or list of (text, {bold,color,size}) tuples."""
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05)
    paras = runs if isinstance(runs, list) else [runs]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        for seg, fmt in ([(para, {})] if isinstance(para, str) else para):
            r = p.add_run()
            r.text = seg
            r.font.name = FONT
            r.font.size = Pt(fmt.get("size", size))
            r.font.bold = fmt.get("bold", bold)
            r.font.color.rgb = fmt.get("color", color)
    return tb


def bullets(slide, x, y, w, h, items, size=18, color=INK, gap=8):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(gap)
        head, _, rest = item.partition("|")
        r = p.add_run()
        r.text = "•  " + head
        r.font.name, r.font.size, r.font.color.rgb = FONT, Pt(size), color
        r.font.bold = bool(rest)
        if rest:
            r2 = p.add_run()
            r2.text = " " + rest
            r2.font.name, r2.font.size, r2.font.color.rgb = FONT, Pt(size), color
    return tb


def header(slide, title, criterion=None, n=None):
    box(slide, 0, 0, W, Inches(0.12), fill=TEAL)
    text(slide, Inches(0.6), Inches(0.38), Inches(9.6), Inches(0.9), title, size=28, color=NAVY, bold=True)
    if criterion:
        tag = box(slide, Inches(10.4), Inches(0.45), Inches(2.4), Inches(0.5), fill=LIGHT, line=TEAL,
                  shape=MSO_SHAPE.ROUNDED_RECTANGLE)
        tf = tag.text_frame
        tf.text = criterion
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        p.runs[0].font.size, p.runs[0].font.bold = Pt(13), True
        p.runs[0].font.color.rgb, p.runs[0].font.name = TEAL, FONT
    if n:
        text(slide, Inches(12.3), Inches(7.0), Inches(0.8), Inches(0.35), str(n), size=11, color=MUTED,
             align=PP_ALIGN.RIGHT)
        text(slide, Inches(0.6), Inches(7.0), Inches(8), Inches(0.35),
             "TatRaksha · Build with AI 2026 · Track 5", size=11, color=MUTED)


def card(slide, x, y, w, h, title, body, accent=TEAL, size=15):
    box(slide, x, y, w, h, fill=LIGHT)
    box(slide, x, y, Inches(0.08), h, fill=accent)
    text(slide, x + Inches(0.25), y + Inches(0.15), w - Inches(0.4), Inches(0.5), title, size=18, color=NAVY,
         bold=True)
    if isinstance(body, list):
        bullets(slide, x + Inches(0.25), y + Inches(0.7), w - Inches(0.4), h - Inches(0.8), body, size=size, gap=4)
    else:
        text(slide, x + Inches(0.25), y + Inches(0.7), w - Inches(0.4), h - Inches(0.8), body, size=size)


def picture(slide, name, x, y, w=None, h=None):
    path = IMG / name
    if path.exists():
        pic = slide.shapes.add_picture(str(path), x, y, width=w, height=h)
        pic.line.color.rgb = RGBColor(0xCB, 0xD5, 0xE1)
        pic.line.width = Pt(1)
        return pic
    box(slide, x, y, w or Inches(6), h or Inches(4), fill=LIGHT)
    return None


def flow(slide, y, steps, colors=None, h=Inches(1.1)):
    n = len(steps)
    gap = Inches(0.35)
    left = Inches(0.6)
    w = Emu(int((W - 2 * left - gap * (n - 1)) / n))
    for i, s in enumerate(steps):
        x = left + i * (w + gap)
        shp = box(slide, x, y, w, h, fill=(colors[i] if colors else NAVY), shape=MSO_SHAPE.ROUNDED_RECTANGLE)
        tf = shp.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = Inches(0.08)
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.text = s
        for p in tf.paragraphs:
            p.alignment = PP_ALIGN.CENTER
            for r in p.runs:
                r.font.size, r.font.bold, r.font.color.rgb, r.font.name = Pt(14 if n <= 5 else 13), True, WHITE, FONT
        if i < n - 1:
            box(slide, x + w + Inches(0.07), y + h / 2 - Inches(0.14), Inches(0.21), Inches(0.28), fill=MUTED,
                shape=MSO_SHAPE.RIGHT_ARROW)


# 1 ── Title ──────────────────────────────────────────────────────────────────────────────────────────
s = prs.slides.add_slide(BLANK)
box(s, 0, 0, W, H, fill=NAVY)
box(s, 0, Inches(5.3), W, Inches(0.08), fill=ORANGE)
text(s, Inches(0.8), Inches(1.4), Inches(11.5), Inches(1.2), "TatRaksha", size=66, color=WHITE, bold=True)
text(s, Inches(0.8), Inches(2.6), Inches(11.5), Inches(1.2),
     "Cyclone anticipatory action & early warning: from an IMD bulletin to village-level action, 72 hours "
     "before landfall", size=26, color=RGBColor(0xCF, 0xE8, 0xEC))
text(s, Inches(0.8), Inches(4.2), Inches(11.5), Inches(0.8),
     "Gemini 3.7 Flash · Earth Engine · Google Maps · Speech & Translation · BigQuery · Cloud Run",
     size=18, color=WHITE)
text(s, Inches(0.8), Inches(5.7), Inches(11.5), Inches(0.9),
     [[("Build with AI hackathon · Track 5", {"bold": True})],
      "Demo: Odisha / Cyclone Fani 2019 and West Bengal / Cyclone Amphan 2020 replays (sample data, labelled)"],
     size=16, color=RGBColor(0xCF, 0xE8, 0xEC))

# 2 ── Problem ────────────────────────────────────────────────────────────────────────────────────────
s = prs.slides.add_slide(BLANK)
header(s, "72 hours, one officer, hundreds of villages", "Problem fit · 20 %", 2)
text(s, Inches(0.6), Inches(1.35), Inches(12), Inches(0.9),
     "“Who do I evacuate, what do I protect, and what do I tell people, in their language, before the "
     "roads flood?”", size=22, color=TEAL, bold=True)
card(s, Inches(0.6), Inches(2.4), Inches(3.9), Inches(3.9), "Forecasts are regional",
     ["IMD bulletins give track, wind and surge guidance for a coastline",
      "Officers must translate them into villages, shelters and substations by hand"], accent=ORANGE)
card(s, Inches(4.7), Inches(2.4), Inches(3.9), Inches(3.9), "Hazards compound",
     ["Surge, rain-flooding and wind hit different villages at different times",
      "Shelters look fine on paper but are cut off when access roads flood"], accent=ORANGE)
card(s, Inches(8.8), Inches(2.4), Inches(3.9), Inches(3.9), "Warnings are slow",
     ["Advisories are drafted manually per area and language",
      "No standard format, no audit trail, money for relief arrives after the event"], accent=ORANGE)
text(s, Inches(0.6), Inches(6.4), Inches(12), Inches(0.5),
     "India's east coast faces 4–5 cyclones a year; the Bay of Bengal has produced most of the deadliest cyclones on record.",
     size=15, color=MUTED)

# 3 ── Solution journey ─────────────────────────────────────────────────────────────────────────────
s = prs.slides.add_slide(BLANK)
header(s, "One journey: bulletin → action", "Problem fit · 20 %", 3)
flow(s, Inches(1.6), ["Gemini parses bulletin; officer verifies track", "Surge + rain-flood + wind swath",
                       "Village Risk + shelter cut-off", "Gemini multimodal situation report",
                       "Advisory → approval → CAP 1.2 → dispatch"],
     colors=[NAVY, TEAL, TEAL, NAVY, ORANGE])
card(s, Inches(0.6), Inches(3.1), Inches(6.0), Inches(3.6), "District Disaster Management Officer",
     ["Replay slider T-72 → T-6 h; every bulletin re-runs the pipeline",
      "Ranked villages with a transparent 30/25/15/15/15 breakdown",
      "Shelters cut off by flooded roads + Google Routes drive time",
      "Situation report with actions per role and deadlines",
      "Odia / Bengali / English advisories with voice, approved before release"])
card(s, Inches(6.8), Inches(3.1), Inches(5.95), Inches(3.6), "State EOC Officer",
     ["Districts and blocks side by side (Odisha, West Bengal)",
      "Parametric trigger probability and payout range (sample policies)",
      "Predicted vs observed flooding for the replay",
      "Advisory, dispatch and audit logs (BigQuery)"], accent=NAVY)

# 4 ── Product screenshot ─────────────────────────────────────────────────────────────────────────────
s = prs.slides.add_slide(BLANK)
header(s, "Who is at risk, and can they reach a shelter?", None, 4)
picture(s, "district-villages.jpg", Inches(0.6), Inches(1.3), h=Inches(5.55))
bullets(s, Inches(9.65), Inches(1.4), Inches(3.3), Inches(5.5), [
    "59,600 people|in High-risk villages at T-24 h (Fani replay)",
    "2 shelters cut off|by the flood-aware road model",
    "Google Routes|3.6 km · ~11 min to the assigned shelter",
    "Every layer labelled|screening estimate, sample data, source + time",
], size=15, gap=10)

# 5 ── AI execution ──────────────────────────────────────────────────────────────────────────────────
s = prs.slides.add_slide(BLANK)
header(s, "Gemini reads and reasons; it never measures",
       "AI & technical · 25 %", 5)
card(s, Inches(0.6), Inches(1.4), Inches(4.0), Inches(2.6), "1 · Bulletin → track",
     "Text/PDF bulletin to schema-validated JSON track; cross-checked by a deterministic parser; officer verifies.",
     size=15)
card(s, Inches(4.75), Inches(1.4), Inches(4.0), Inches(2.6), "2 · Multimodal sitrep",
     "Hazard-map image + bulletin + exposure tables in one call → headline, key numbers, actions by role with "
     "deadlines, uncertainties.", size=15)
card(s, Inches(8.9), Inches(1.4), Inches(3.85), Inches(2.6), "3 · Advisory drafting",
     "English drafts with six mandatory fields; rejected if a field is missing. Odia/Bengali from approved "
     "templates.", size=15)
card(s, Inches(0.6), Inches(4.2), Inches(12.15), Inches(2.6), "Trust by construction", [
    "Observed data → model estimate → AI interpretation → recommendation, kept separate in the data and UI",
    "Every number in the report is checked against the pipeline output; ungrounded items are removed and listed",
    "Each AI record stores model_name, model_version, prompt_version (versioned prompts in ai/prompts/)",
    "Rate limits or failures fall back to a labelled deterministic path; the live/demo badge is verified by real calls",
], accent=NAVY)

# 6 ── Sitrep screenshot + Google stack ─────────────────────────────────────────────────────────────
s = prs.slides.add_slide(BLANK)
header(s, "Google AI & Cloud stack, live in the prototype", "AI & technical · 25 %", 6)
picture(s, "district-sitrep.jpg", Inches(0.6), Inches(1.3), h=Inches(5.5))
rows = [("Gemini 3.7 Flash (Vertex AI)", "parse · sitrep · drafts"),
        ("Maps JS · Routes · Geocoding", "basemap · shelter drive time"),
        ("Text-to-Speech · Speech-to-Text", "voice advisory + read-back"),
        ("Cloud Translation", "back-translation for review"),
        ("BigQuery", "audit + dispatch log"),
        ("Cloud Storage", "map renders · CAP · audio"),
        ("Cloud Run · Secret Manager", "API + 2 dashboards"),
        ("Earth Engine", "DEM, water, rain, WorldPop (ready)")]
y = Inches(1.35)
for name, role in rows:
    text(s, Inches(9.55), y, Inches(3.6), Inches(0.65),
         [[(name, {"bold": True, "color": NAVY, "size": 14})], [(role, {"color": MUTED, "size": 12})]], spacing=1.0)
    y += Inches(0.68)

# 7 ── Hazard & risk science ────────────────────────────────────────────────────────────────────────
s = prs.slides.add_slide(BLANK)
header(s, "Transparent hazard models an officer can explain", "AI & technical · 25 %", 7)
card(s, Inches(0.6), Inches(1.4), Inches(6.0), Inches(2.5), "Storm surge (screening)", [
    "Inverse barometer + wind set-up × shelf factor, right-of-track peak",
    "Bathtub fill connected to the sea; expected vs high case",
    "Labelled: defer to official IMD / INCOIS guidance"], size=15)
card(s, Inches(6.8), Inches(1.4), Inches(5.95), Inches(2.5), "Rain-flood + wind", [
    "Rain-flood likelihood from 72 h rain, low-lying terrain, drainage",
    "Modified Rankine wind swath, gale arrival time per village",
    "Validated against observed flooding in each replay"], size=15)
card(s, Inches(0.6), Inches(4.1), Inches(12.15), Inches(2.7), "Village Risk score (0–100)", [
    "Surge 30 · Rainfall flood 25 · Wind 15 · Population 15 · Vulnerability 15, weights set per state in its adapter",
    "Road-network cut-off: a shelter counts only if it is dry and reachable on roads that stay open",
    "Earth Engine path (Copernicus DEM, JRC water, GPM IMERG, WorldPop, Open Buildings) swaps in real layers",
], accent=ORANGE)

# 8 ── Advisory & languages ──────────────────────────────────────────────────────────────────────────
s = prs.slides.add_slide(BLANK)
header(s, "Warnings people understand, released by a human", "Impact · 15 %", 8)
flow(s, Inches(1.5), ["Draft per village, audience, stage", "Back-translation + review", "Officer approval",
                       "CAP 1.2 (validated)", "SMS · voice · feed (sandbox)"],
     colors=[TEAL, TEAL, ORANGE, NAVY, NAVY])
card(s, Inches(0.6), Inches(3.0), Inches(6.0), Inches(3.7), "Languages and voice", [
    "English, Odia, Bengali; mandatory fields preserved (what, where, when, action, shelter, authority)",
    "Cloud Text-to-Speech voice advisory; Speech-to-Text reads it back",
    "Adding Telugu or Tamil is a template + config change"], size=15)
card(s, Inches(6.8), Inches(3.0), Inches(5.95), Inches(3.7), "Standards and accountability", [
    "CAP 1.2: the format of India's national alert gateway (NDMA)",
    "Consequential actions need human approval; two roles only",
    "Every draft, approval and dispatch in an append-only audit log (BigQuery)"], accent=NAVY, size=15)

# 9 ── Depth & reach ────────────────────────────────────────────────────────────────────────────────
s = prs.slides.add_slide(BLANK)
header(s, "Built for every coastal state, not one district", "Depth & reach · 20 %", 9)
picture(s, "state-eoc.jpg", Inches(6.9), Inches(1.35), w=Inches(5.85))
bullets(s, Inches(0.6), Inches(1.45), Inches(6.0), Inches(5.4), [
    "Canonical disaster-risk schema|(villages, shelters, assets, roads, hazards, advisories)",
    "State adapters in JSON|districts, languages, authority, CAP sender, risk weights, field mapping",
    "Two states live|Odisha (Fani 2019), West Bengal (Amphan 2020) on the same services",
    "Global layers|Earth Engine datasets cover every Indian coast out of the box",
    "Onboarding a state|adapter + data files + approved templates, no code fork",
], size=17, gap=12)

# 10 ── Deployability & scale ───────────────────────────────────────────────────────────────────────
s = prs.slides.add_slide(BLANK)
header(s, "Cloud-native, API-first, pilot-ready", "Deployability · 20 %", 10)
flow(s, Inches(1.45), ["One village", "One district", "One state", "All coastal states", "Bay of Bengal region",
                        "Regional network"], colors=[TEAL, TEAL, TEAL, NAVY, NAVY, ORANGE], h=Inches(0.95))
card(s, Inches(0.6), Inches(2.75), Inches(6.0), Inches(4.0), "Deployable today", [
    "Cloud Run (asia-south1): FastAPI + two Next.js apps, one repeatable deploy script",
    "Secrets in Secret Manager; service-account auth; public repo, no keys in git",
    "Every integration degrades to a labelled demo mode, so it runs anywhere",
    "REST API (PRD §37) + CAP feed for state alert gateways"], size=15)
card(s, Inches(6.8), Inches(2.75), Inches(5.95), Inches(4.0), "Pilot path", [
    "Shadow mode next cyclone season alongside the OSDMA / WBDMD process",
    "Compare Village Risk and cut-offs with observed flooding (built-in validation)",
    "Firestore + Earth Engine switch on by configuration",
    "Vertex AI flood model trained on Sentinel-1 flood extents"], accent=ORANGE, size=15)

# 11 ── Impact + cross-border ─────────────────────────────────────────────────────────────────────────
s = prs.slides.add_slide(BLANK)
header(s, "Earlier, targeted evacuation and faster money", "Impact · 15 %", 11)
card(s, Inches(0.6), Inches(1.4), Inches(4.0), Inches(3.0), "Lives",
     "Evacuation orders reach the highest-risk villages first, with a reachable shelter named in the "
     "advisory and in the local language.", size=15)
card(s, Inches(4.75), Inches(1.4), Inches(4.0), Inches(3.0), "Infrastructure",
     "Exposed substations, hospitals and roads flagged 24–72 h ahead so utilities can harden and "
     "pre-position crews.", size=15)
card(s, Inches(8.9), Inches(1.4), Inches(3.85), Inches(3.0), "Liquidity",
     "Parametric trigger estimates let insurers and the state prepare payouts before landfall.", size=15)
card(s, Inches(0.6), Inches(4.6), Inches(12.15), Inches(2.2), "Cross-border and BRICS portability", [
    "Bangladesh and Myanmar share the Bay of Bengal cyclones: Bangla and Burmese are configuration, CAP is international",
    "State adapter → country adapter: the same canonical model serves any national met service and alert gateway",
], accent=NAVY, size=15)

# 12 ── Status & ask ────────────────────────────────────────────────────────────────────────────────
s = prs.slides.add_slide(BLANK)
header(s, "Where we are, honestly", None, 12)
card(s, Inches(0.6), Inches(1.4), Inches(6.0), Inches(5.3), "Working now", [
    "Full journey: bulletin → hazards → village risk → Gemini sitrep → approved CAP advisory",
    "Live: Gemini 3.7 Flash, Maps Routes/Geocoding, Translation, TTS, STT, BigQuery, Cloud Storage",
    "Two states, three languages, State EOC with parametric and validation",
    "40 API tests incl. an end-to-end journey; deployable Docker images"], size=15)
card(s, Inches(6.8), Inches(1.4), Inches(5.95), Inches(5.3), "Next", [
    "Earth Engine registration → real DEM, rainfall and population layers",
    "Vertex AI flood model; live IMD forecast ingest",
    "Native-speaker review of Odia/Bengali templates; Odia voice",
    "Andhra Pradesh rain-flood scenario; SMS/voice sandbox providers",
    "Shadow-mode pilot with one coastal district"], accent=ORANGE, size=15)

prs.save(OUT)
print(f"wrote {OUT.relative_to(HERE.parents[1])} ({len(prs.slides)} slides)")
